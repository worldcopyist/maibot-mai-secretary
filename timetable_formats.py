"""Decode offline timetable exports into bounded two-dimensional tables.

File extensions are only hints. Binary signatures and table structure decide how
the data is read; this module never opens links or evaluates workbook formulas.
"""

from __future__ import annotations

import base64
import csv
import io
import zipfile
from datetime import date, datetime
from html.parser import HTMLParser
from typing import Any

FILE_PREFIX = "TIMETABLE-FILE:"
LEGACY_XLS_PREFIX = "XLS-BASE64:"
MAX_FILE_BYTES = 48_000
MAX_ROWS = 1_000
MAX_COLUMNS = 40
MAX_CELLS = 20_000


class _TablesFromHtml(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self.table_depth = 0
        self.rows: list[list[str]] = []
        self.row: list[str] | None = None
        self.cell: list[str] | None = None
        self.cell_span = 1
        self.skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self.skip_depth += 1
            return
        if self.skip_depth:
            return
        if tag == "table":
            if self.table_depth == 0:
                self.rows = []
            self.table_depth += 1
        elif tag == "tr" and self.table_depth == 1:
            self.row = []
        elif tag in {"td", "th"} and self.table_depth == 1 and self.row is not None:
            self.cell = []
            attributes = dict(attrs)
            try:
                self.cell_span = min(max(int(attributes.get("colspan") or 1), 1), MAX_COLUMNS)
            except ValueError:
                self.cell_span = 1
        elif tag == "br" and self.cell is not None:
            self.cell.append("\n")
        elif tag in {"div", "p"} and self.cell is not None and self.cell:
            self.cell.append("\n\n")
        elif tag == "font" and self.cell is not None:
            title = dict(attrs).get("title") or ""
            if "周" in title and len(title) < 500:
                self.cell.extend(("\n", title))

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self.skip_depth = max(0, self.skip_depth - 1)
            return
        if self.skip_depth:
            return
        if tag in {"td", "th"} and self.cell is not None and self.row is not None:
            value = "".join(self.cell).strip()
            self.row.extend([value, *([""] * (self.cell_span - 1))])
            self.cell = None
        elif tag == "tr" and self.row is not None:
            if any(value for value in self.row):
                self.rows.append(self.row)
            self.row = None
        elif tag == "table" and self.table_depth:
            self.table_depth -= 1
            if self.table_depth == 0 and self.rows:
                self.tables.append(self.rows)
                self.rows = []

    def handle_data(self, data: str) -> None:
        if self.cell is not None and not self.skip_depth:
            self.cell.append(data)


def _decode_text(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("课表文本编码无法识别；请另存为 UTF-8")


def _read_xls(payload: bytes) -> list[list[list[str]]]:
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError("服务器缺少 xlrd 依赖，无法读取 .xls 文件") from exc
    try:
        book = xlrd.open_workbook(file_contents=payload)
        def cell_text(sheet: Any, row: int, col: int) -> str:
            cell = sheet.cell(row, col)
            if cell.ctype == xlrd.XL_CELL_DATE:
                return xlrd.xldate_as_datetime(cell.value, book.datemode).date().isoformat()
            if cell.ctype == xlrd.XL_CELL_NUMBER and cell.value.is_integer():
                return str(int(cell.value))
            return str(cell.value).strip() if cell.value is not None else ""
        return [
            [[cell_text(sheet, row, col) for col in range(min(sheet.ncols, MAX_COLUMNS))]
             for row in range(min(sheet.nrows, MAX_ROWS))]
            for sheet in book.sheets()[:5]
        ]
    except Exception as exc:
        raise ValueError(f".xls 文件无法解析：{exc}") from exc


def _read_xlsx(payload: bytes) -> list[list[list[str]]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("服务器缺少 openpyxl 依赖，无法读取 .xlsx 文件") from exc
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            infos = archive.infolist()
            if len(infos) > 200 or sum(item.file_size for item in infos) > 10_000_000:
                raise ValueError(".xlsx 解压内容过大")
        book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
        try:
            tables = []
            for sheet in book.worksheets[:5]:
                rows = []
                for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or MAX_ROWS, MAX_ROWS),
                                           max_col=min(sheet.max_column or MAX_COLUMNS, MAX_COLUMNS), values_only=True):
                    rows.append([value.date().isoformat() if isinstance(value, datetime) else
                                 value.isoformat() if isinstance(value, date) else
                                 str(value).strip() if value is not None else "" for value in row])
                tables.append(rows)
            return tables
        finally:
            book.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f".xlsx 文件无法解析：{exc}") from exc


def read_tables(content: str) -> tuple[list[list[list[str]]], str]:
    """Read a browser-uploaded file without trusting its extension alone."""
    if content.startswith(LEGACY_XLS_PREFIX):
        suffix = "xls"
        encoded = content[len(LEGACY_XLS_PREFIX):]
    elif content.startswith(FILE_PREFIX):
        try:
            suffix, encoded = content[len(FILE_PREFIX):].split(":", 1)
        except ValueError as exc:
            raise ValueError("课表上传数据格式无效") from exc
        suffix = suffix.lower().strip(".")
    else:
        raise ValueError("请从课表页面上传 .xls、.xlsx、.csv、.tsv 或 .html 文件")
    if suffix not in {"xls", "xlsx", "csv", "tsv", "html", "htm"}:
        raise ValueError("不支持的课表文件格式")
    try:
        payload = base64.b64decode(encoded.strip(), validate=True)
    except ValueError as exc:
        raise ValueError("课表文件编码无效") from exc
    if not payload or len(payload) > MAX_FILE_BYTES:
        raise ValueError("课表文件为空或超过 48 KB 上限")
    if payload.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return _validate(_read_xls(payload)), "课表 Excel 文件"
    if payload.startswith(b"PK\x03\x04"):
        return _validate(_read_xlsx(payload)), "课表 XLSX 文件"
    text = _decode_text(payload)
    if "<table" in text.lower():
        parser = _TablesFromHtml()
        parser.feed(text)
        if not parser.tables:
            raise ValueError("HTML 文件中没有可读取的表格")
        return _validate(parser.tables), "课表 HTML 文件"
    if suffix not in {"csv", "tsv"}:
        raise ValueError("文件内容不是可识别的 Excel、HTML 或 CSV 课表")
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",\t;").delimiter
    except csv.Error:
        delimiter = "\t" if suffix == "tsv" else ","
    rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return _validate([rows]), "课表 CSV 文件"


def _validate(tables: list[list[list[Any]]]) -> list[list[list[str]]]:
    if len(tables) > 20 or sum(len(table) for table in tables) > MAX_ROWS * 5:
        raise ValueError("课表工作表或行数过多")
    if sum(len(row) for table in tables for row in table) > MAX_CELLS:
        raise ValueError("课表单元格过多")
    return tables
